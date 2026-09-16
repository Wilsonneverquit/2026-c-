
from pathlib import Path
import numpy as np
from openpyxl import load_workbook, Workbook
from scipy.optimize import milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parent
DATA = ROOT.parent / "附件" / "附件1.xlsx"
TEMPLATE = ROOT.parent / "附件" / "附件5" / "result1.xlsx"
OUTPUT = ROOT / "result1.xlsx"

DT = 1.0 / 6.0
ETA_C = 0.9
ETA_D = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
P_MAX = 5000.0
E_INITIAL = 6000.0


def read_source():
    workbook = load_workbook(DATA, read_only=True, data_only=True)
    rows = list(workbook.active.values)[1:]
    workbook.close()
    times = [row[1] for row in rows]
    price = np.asarray([row[2] for row in rows], dtype=float)
    load = np.asarray([row[3] for row in rows], dtype=float)
    pv = np.asarray([row[4] for row in rows], dtype=float)
    if len(rows) != 144:
        raise ValueError(f"附件1应有144个时段，实际为{len(rows)}")
    return times, price, load, pv


def solve(price, load, pv):
    n = len(price)
    # x=[grid, charge, discharge, curtailment, E(0..n), binary u]
    grid = 0
    charge = n
    discharge = 2 * n
    curtail = 3 * n
    energy = 4 * n
    binary = 5 * n + 1
    nvar = binary + n

    objective = np.zeros(nvar)
    objective[grid:grid + n] = price * DT
    lower = np.zeros(nvar)
    upper = np.full(nvar, np.inf)
    upper[charge:charge + n] = P_MAX
    upper[discharge:discharge + n] = P_MAX
    upper[curtail:curtail + n] = pv
    lower[energy:energy + n + 1] = E_MIN
    upper[energy:energy + n + 1] = E_MAX
    lower[energy] = upper[energy] = E_INITIAL
    lower[energy + n] = upper[energy + n] = E_INITIAL
    upper[binary:] = 1.0
    integrality = np.zeros(nvar)
    integrality[binary:] = 1

    matrix = lil_matrix((4 * n + 1, nvar))
    row_lower = np.full(4 * n + 1, -np.inf)
    row_upper = np.full(4 * n + 1, np.inf)
    for t in range(n):
        # grid + discharge + pv - curtail = load + charge
        matrix[t, grid + t] = 1
        matrix[t, charge + t] = -1
        matrix[t, discharge + t] = 1
        matrix[t, curtail + t] = -1
        row_lower[t] = row_upper[t] = load[t] - pv[t]

        # E[t+1]-E[t]-eta_c*charge*dt+discharge*dt/eta_d=0
        dyn = n + t
        matrix[dyn, energy + t + 1] = 1
        matrix[dyn, energy + t] = -1
        matrix[dyn, charge + t] = -ETA_C * DT
        matrix[dyn, discharge + t] = DT / ETA_D
        row_lower[dyn] = row_upper[dyn] = 0

        # charge <= P_MAX*u; discharge <= P_MAX*(1-u)
        mutex_charge = 2 * n + t
        matrix[mutex_charge, charge + t] = 1
        matrix[mutex_charge, binary + t] = -P_MAX
        row_upper[mutex_charge] = 0
        mutex_discharge = 3 * n + t
        matrix[mutex_discharge, discharge + t] = 1
        matrix[mutex_discharge, binary + t] = P_MAX
        row_upper[mutex_discharge] = P_MAX

    matrix[4 * n, energy] = 1
    matrix[4 * n, energy + n] = -1
    row_lower[4 * n] = row_upper[4 * n] = 0

    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(matrix.tocsr(), row_lower, row_upper),
        options={"mip_rel_gap": 1e-9},
    )
    if result.status != 0 or result.x is None:
        raise RuntimeError(f"求解失败: status={result.status}, message={result.message}")
    x = result.x
    return {
        "grid": x[grid:grid + n],
        "charge": x[charge:charge + n],
        "discharge": x[discharge:discharge + n],
        "curtail": x[curtail:curtail + n],
        "energy": x[energy:energy + n + 1],
        "cost": float(result.fun),
    }


def template_labels():
    workbook = load_workbook(TEMPLATE, read_only=True, data_only=True)
    sheet = workbook["计划购电量"]
    labels = [sheet.cell(row=i, column=1).value for i in range(2, 146)]
    workbook.close()
    return labels


def write_result(times, price, load, pv, result):
    labels = template_labels()
    workbook = Workbook()
    plan = workbook.active
    plan.title = "计划购电量"
    plan.append(["时间段", "购电量"])
    grid_kwh = result["grid"] * DT
    for label, value in zip(labels, grid_kwh):
        plan.append([label, float(value)])

    storage = workbook.create_sheet("充放电量")
    storage.append(["时间段", "充电量", "放电量", "时刻", "储电量"])
    for segment in range(6):
        start, end = segment * 24, (segment + 1) * 24
        storage.append([
            f"{segment * 4}:00-{(segment + 1) * 4}:00",
            float(np.sum(result["charge"][start:end]) * DT),
            float(np.sum(result["discharge"][start:end]) * DT),
            "0:00" if segment == 0 else ("24:00" if segment == 1 else None),
            float(result["energy"][0] if segment == 0 else result["energy"][-1]) if segment < 2 else None,
        ])
    workbook.save(OUTPUT)


def main():
    times, price, load, pv = read_source()
    result = solve(price, load, pv)
    write_result(times, price, load, pv, result)
    grid, charge, discharge, curtail = result["grid"], result["charge"], result["discharge"], result["curtail"]
    balance = grid + discharge + pv - curtail - load - charge
    print(f"输出文件: {OUTPUT}")
    print(f"最优购电费: {result['cost']:.6f} 元")
    print(f"全天购电量: {np.sum(grid) * DT:.6f} kWh")
    print(f"充电量: {np.sum(charge) * DT:.6f} kWh")
    print(f"放电量: {np.sum(discharge) * DT:.6f} kWh")
    print(f"弃光量: {np.sum(curtail) * DT:.6f} kWh")
    print(f"功率平衡最大误差: {np.max(np.abs(balance)):.3e} kW")
    print(f"储能范围: {result['energy'].min():.6f}—{result['energy'].max():.6f} kWh")
    print(f"首末储能: {result['energy'][0]:.6f}, {result['energy'][-1]:.6f} kWh")
    print(f"模板时段首尾: {template_labels()[0]} / {template_labels()[-1]}")


if __name__ == "__main__":
    main()
