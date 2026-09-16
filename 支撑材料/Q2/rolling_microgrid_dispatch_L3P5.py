"""
问题二（优化窗口版）：负荷3天分组滚动预测 + 光伏5天滚动预测 + 日前优化 + 日内执行。

相对旧方案（方案D）的唯一变化是预测窗口长度：
  - 负荷：星期分组（周五/周六为一组，周一~周四+周日为另一组），取同组最近 3 个历史日；
  - 光伏：取最近 5 个自然日；
  - 两者均逐十分钟时点取中位数，再做三点移动平滑、截断为非负。
规划模型、日内执行、储能参数与方案D完全一致。

数据口径
--------
附件 1、附件 2 每天均有 144 个十分钟时点，按附件 5 的模板顺序解释：
0:10 -> 0:10-0:20，...，23:50 -> 23:50-次日0:00，
0:00+1 -> 次日0:00-0:10。脚本只做顺序对应，不整体平移数据。

变量在优化模型中使用功率 kW；写入 Excel 的购电、充电、放电均为
十分钟电量 kWh，因此统一乘 DT = 1/6 h。

运行示例
--------
python rolling_microgrid_dispatch_L3P5.py
python rolling_microgrid_dispatch_L3P5.py --pseudo-test
python rolling_microgrid_dispatch_L3P5.py --end-date 2025-02-03
"""

from __future__ import annotations

import argparse
import copy
import math
import os
import shutil
from dataclasses import dataclass
from datetime import time as dt_time
from pathlib import Path

import cvxpy as cp
os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parent / ".mplconfig"))
import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

matplotlib.use("Agg")
import matplotlib.pyplot as plt


# ------------------------------ 路径与常量 ------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent
C_PROBLEM_DIR = SCRIPT_DIR.parent
ATTACHMENT_DIR = C_PROBLEM_DIR / "附件"
PRICE_FILE = ATTACHMENT_DIR / "附件1.xlsx"
ACTUAL_FILE = ATTACHMENT_DIR / "附件2.xlsx"
TEMPLATE_FILE = ATTACHMENT_DIR / "附件5" / "result2.xlsx"

# 旧方案结果（用于对比）：方案D（负荷7天/光伏7天）与最初基线（7自然日/不分组）
OLD_SCHEME_DAILY = C_PROBLEM_DIR / "对比基线" / "每日费用_D.csv"
BASE_SCHEME_DAILY = C_PROBLEM_DIR / "对比基线" / "每日费用_基线.csv"

DT = 1.0 / 6.0                    # 每个时段为 10 min = 1/6 h
N_SLOTS = 144
E_INIT = 6000.0                   # 2025-01-01 0:00 初始储能，kWh
E_MIN = 1200.0                    # 允许运行下限，kWh
E_MAX = 10800.0                   # 允许运行上限，kWh
P_MAX = 5000.0                    # 最大充/放电功率，kW
ETA_C = 0.9                       # 充电效率
ETA_D = 0.9                       # 放电效率
PLAN_TERMINAL_SOC = 6000.0        # 每个日前模型的计划终端约束，kWh
EMERGENCY_PRICE_FACTOR = 5.0
TOL = 1e-5

# 阶梯测试得到的最优预测窗口
LOAD_WINDOW = 3                   # 负荷：同星期组最近 3 个历史日
PV_WINDOW = 5                     # 光伏：最近 5 个自然日

START_DATE = pd.Timestamp("2025-02-01")
END_DATE = pd.Timestamp("2025-12-31")
SPECIAL_DATES = {
    pd.Timestamp("2025-03-20"),
    pd.Timestamp("2025-06-21"),
    pd.Timestamp("2025-09-23"),
    pd.Timestamp("2025-12-21"),
}


@dataclass
class DayAheadResult:
    """日前 MILP 的计划解，全部功率变量单位均为 kW。"""

    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    soc: np.ndarray
    mode: np.ndarray
    objective: float
    status: str


@dataclass
class IntradayResult:
    """日内实际执行结果；soc 有 145 个边界时刻，其余数组有 144 个时段。"""

    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    curtail: np.ndarray
    grid_spill: np.ndarray
    soc: np.ndarray


def moving_average_3(values: np.ndarray) -> np.ndarray:
    """对单日 144 点序列做三点中心移动平均，边界使用端点复制。"""

    padded = np.pad(np.asarray(values, dtype=float), (1, 1), mode="edge")
    return np.convolve(padded, np.ones(3) / 3.0, mode="valid")


def forecast_from_history(
    load_history: np.ndarray,
    pv_history: np.ndarray,
    target_date: pd.Timestamp | None = None,
    history_dates: pd.DatetimeIndex | None = None,
) -> tuple[np.ndarray, np.ndarray]:
    """
    负荷预测：从目标日所属星期组中选最近 LOAD_WINDOW 个历史日，逐时点取中位数。

    星期组A为周五、周六；星期组B为周一、周二、周三、周四、周日。
    光伏没有从热力图得到相同的星期规律，因此使用最近 PV_WINDOW 个自然日中位数。
    两类预测最后均做三点移动平滑。

    参数中的历史数据只能包含预测日之前已经观测到的实际值，从而严格避免
    使用未来信息。2025 年 1 月构成初始训练集；进入 2 月后每天追加当天实际值。
    """

    if len(load_history) < LOAD_WINDOW or len(pv_history) < PV_WINDOW:
        raise ValueError("预测历史实际数据不足。")
    if target_date is not None and history_dates is not None:
        if len(history_dates) != len(load_history):
            raise ValueError("历史日期数量与负荷历史天数不一致。")
        target_in_friday_saturday = target_date.weekday() in {4, 5}
        eligible = np.asarray([
            (date.weekday() in {4, 5}) == target_in_friday_saturday
            for date in history_dates
        ])
        selected = np.flatnonzero(eligible)[-LOAD_WINDOW:]
        if len(selected) < LOAD_WINDOW:
            raise ValueError(f"{target_date.date()} 所属星期组不足{LOAD_WINDOW}个历史样本。")
        load_median = np.median(load_history[selected], axis=0)
    else:
        load_median = np.median(load_history[-LOAD_WINDOW:], axis=0)
    pv_median = np.median(pv_history[-PV_WINDOW:], axis=0)
    load_forecast = np.clip(moving_average_3(load_median), 0.0, None)
    pv_forecast = np.clip(moving_average_3(pv_median), 0.0, None)
    return load_forecast, pv_forecast


def read_input_data() -> tuple[np.ndarray, pd.DatetimeIndex, np.ndarray, np.ndarray, list[str]]:
    """读取电价、全年负荷、全年光伏实际功率及模板的 144 个时段标题。"""

    price_frame = pd.read_excel(PRICE_FILE)
    if price_frame.shape[0] != N_SLOTS:
        raise ValueError(f"附件1应有 {N_SLOTS} 行，实际为 {price_frame.shape[0]} 行。")
    price = price_frame.iloc[:, 2].to_numpy(dtype=float)

    load_frame = pd.read_excel(ACTUAL_FILE, sheet_name=0)
    pv_frame = pd.read_excel(ACTUAL_FILE, sheet_name=1)
    dates = pd.DatetimeIndex(pd.to_datetime(load_frame.iloc[:, 0])).normalize()
    pv_dates = pd.DatetimeIndex(pd.to_datetime(pv_frame.iloc[:, 0])).normalize()
    if not dates.equals(pv_dates):
        raise ValueError("附件2的负荷日期与光伏日期不一致。")
    load_actual = load_frame.iloc[:, 1:145].to_numpy(dtype=float)
    pv_actual = pv_frame.iloc[:, 1:145].to_numpy(dtype=float)
    if load_actual.shape != (365, N_SLOTS) or pv_actual.shape != (365, N_SLOTS):
        raise ValueError("附件2必须包含365天、每天144个十分钟功率值。")

    template_book = load_workbook(TEMPLATE_FILE, read_only=True, data_only=False)
    plan_sheet = template_book["计划购电量"]
    slot_labels = [plan_sheet.cell(1, column).value for column in range(2, 146)]
    template_book.close()
    return price, dates, load_actual, pv_actual, slot_labels


def build_day_ahead_problem(price: np.ndarray) -> tuple[cp.Problem, dict[str, object]]:
    """
    构造可重复使用的单日 MILP。

    二进制变量 mode[t] = 1 表示允许充电，= 0 表示允许放电，因此严格禁止
    同一时段同时充放电。日前模型仅使用预测数据，紧急购电不进入该模型。
    """

    load_parameter = cp.Parameter(N_SLOTS, nonneg=True, name="load_forecast")
    pv_parameter = cp.Parameter(N_SLOTS, nonneg=True, name="pv_forecast")
    initial_soc_parameter = cp.Parameter(nonneg=True, name="initial_soc")

    grid = cp.Variable(N_SLOTS, nonneg=True, name="grid_plan")
    charge = cp.Variable(N_SLOTS, nonneg=True, name="charge_plan")
    discharge = cp.Variable(N_SLOTS, nonneg=True, name="discharge_plan")
    pv_used = cp.Variable(N_SLOTS, nonneg=True, name="pv_used_plan")
    soc = cp.Variable(N_SLOTS + 1, name="soc_plan")
    mode = cp.Variable(N_SLOTS, boolean=True, name="charge_mode")

    constraints = [
        soc[0] == initial_soc_parameter,
        soc[N_SLOTS] == PLAN_TERMINAL_SOC,
        soc >= E_MIN,
        soc <= E_MAX,
        charge <= P_MAX * mode,
        discharge <= P_MAX * (1.0 - mode),
        pv_used <= pv_parameter,
        grid + discharge + pv_used == load_parameter + charge,
        soc[1:] == soc[:-1] + ETA_C * charge * DT - discharge * DT / ETA_D,
    ]
    objective = cp.Minimize(cp.sum(cp.multiply(price * DT, grid)))
    problem = cp.Problem(objective, constraints)
    objects: dict[str, object] = {
        "load_parameter": load_parameter,
        "pv_parameter": pv_parameter,
        "initial_soc_parameter": initial_soc_parameter,
        "grid": grid,
        "charge": charge,
        "discharge": discharge,
        "pv_used": pv_used,
        "soc": soc,
        "mode": mode,
    }
    return problem, objects


def solve_day_ahead(
    problem: cp.Problem,
    objects: dict[str, object],
    load_forecast: np.ndarray,
    pv_forecast: np.ndarray,
    initial_soc: float,
) -> DayAheadResult:
    """把当天预测和真实日初 SOC 代入 MILP，并调用 SCIP 求解。"""

    objects["load_parameter"].value = load_forecast
    objects["pv_parameter"].value = pv_forecast
    objects["initial_soc_parameter"].value = float(initial_soc)
    objective = problem.solve(
        solver=cp.SCIP,
        verbose=False,
        scip_params={"limits/gap": 1e-7, "display/verblevel": 0},
    )
    if problem.status not in {cp.OPTIMAL, cp.OPTIMAL_INACCURATE}:
        raise RuntimeError(f"日前 MILP 求解失败，状态为 {problem.status}。")

    def clean(name: str) -> np.ndarray:
        values = np.asarray(objects[name].value, dtype=float)
        values[np.abs(values) < 1e-8] = 0.0
        return values

    return DayAheadResult(
        grid=clean("grid"),
        charge=clean("charge"),
        discharge=clean("discharge"),
        pv_used=clean("pv_used"),
        soc=clean("soc"),
        mode=np.rint(clean("mode")),
        objective=float(objective),
        status=problem.status,
    )


def execute_intraday(
    grid_plan: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
    initial_soc: float,
) -> IntradayResult:
    """
    按十分钟顺序执行固定的正常购电计划，并实时调整储能。

    规则如下：
    1. 正常购电严格取 grid_plan，不允许日内更改。
    2. 当正常购电与真实光伏不足以满足真实负荷时，先在功率/SOC范围内放电，
       剩余缺口使用 5 倍电价紧急购电。
    3. 当供应超过负荷时，优先在功率/SOC范围内充电。
    4. 真实光伏在满足负荷后仍未能用于充电的部分才计为弃光；由日前购电
       预测偏高造成的剩余单独计为 grid_spill，绝不混入弃光。

    该策略只用当前时段已揭示的真实数据，属于可执行的因果在线策略。
    """

    charge = np.zeros(N_SLOTS)
    discharge = np.zeros(N_SLOTS)
    emergency = np.zeros(N_SLOTS)
    curtail = np.zeros(N_SLOTS)
    grid_spill = np.zeros(N_SLOTS)
    soc = np.zeros(N_SLOTS + 1)
    soc[0] = float(initial_soc)

    for slot in range(N_SLOTS):
        supply_minus_load = grid_plan[slot] + pv_actual[slot] - load_actual[slot]
        if supply_minus_load >= 0.0:
            capacity_limited_charge = max(0.0, (E_MAX - soc[slot]) / (ETA_C * DT))
            charge[slot] = min(supply_minus_load, P_MAX, capacity_limited_charge)

            # 先把充电量归因于真实光伏盈余，再归因于计划购电盈余。
            true_pv_surplus = max(pv_actual[slot] - load_actual[slot], 0.0)
            pv_charge = min(charge[slot], true_pv_surplus)
            curtail[slot] = max(true_pv_surplus - pv_charge, 0.0)
            residual_surplus = supply_minus_load - charge[slot]
            grid_spill[slot] = max(residual_surplus - curtail[slot], 0.0)
        else:
            deficit = -supply_minus_load
            energy_limited_discharge = max(0.0, (soc[slot] - E_MIN) * ETA_D / DT)
            discharge[slot] = min(deficit, P_MAX, energy_limited_discharge)
            emergency[slot] = max(deficit - discharge[slot], 0.0)

        soc[slot + 1] = (
            soc[slot]
            + ETA_C * charge[slot] * DT
            - discharge[slot] * DT / ETA_D
        )
        # 仅消除浮点误差，不改变调度决策。
        if abs(soc[slot + 1] - E_MIN) < 1e-8:
            soc[slot + 1] = E_MIN
        if abs(soc[slot + 1] - E_MAX) < 1e-8:
            soc[slot + 1] = E_MAX

    return IntradayResult(charge, discharge, emergency, curtail, grid_spill, soc)


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    """复制附件5模板行的样式，以便扩展到完整334天。"""

    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy.copy(source.alignment)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.protection = copy.copy(source.protection)


def write_result_workbook(
    output_path: Path,
    dates: pd.DatetimeIndex,
    slot_labels: list[str],
    price: np.ndarray,
    grid_plan: np.ndarray,
    actual_charge: np.ndarray,
    actual_discharge: np.ndarray,
    emergency: np.ndarray,
    soc: np.ndarray,
) -> None:
    """按附件5模板写出 result2.xlsx：只写数值，完整保留模板样式与日期格式(mm-dd-yy)。"""

    target = Path(output_path)
    work = target.with_name(target.stem + "__tmp" + target.suffix)
    shutil.copyfile(TEMPLATE_FILE, work)
    workbook = load_workbook(work)
    plan_sheet = workbook["计划购电量"]
    storage_sheet = workbook["充放电量"]
    emergency_sheet = workbook["紧急购电量"]

    def snapshot(sheet, row: int, ncols: int):
        return [{
            "font": copy.copy(sheet.cell(row, c).font),
            "fill": copy.copy(sheet.cell(row, c).fill),
            "border": copy.copy(sheet.cell(row, c).border),
            "alignment": copy.copy(sheet.cell(row, c).alignment),
            "protection": copy.copy(sheet.cell(row, c).protection),
            "number_format": sheet.cell(row, c).number_format,
        } for c in range(1, ncols + 1)]

    def apply(sheet, row: int, style) -> None:
        for c, st in enumerate(style, start=1):
            cell = sheet.cell(row, c)
            cell.font = copy.copy(st["font"])
            cell.fill = copy.copy(st["fill"])
            cell.border = copy.copy(st["border"])
            cell.alignment = copy.copy(st["alignment"])
            cell.protection = copy.copy(st["protection"])
            cell.number_format = st["number_format"]

    # 1) 计划购电量：模板已预留334天数据行且样式正确，只覆盖数值，日期格式沿用模板 mm-dd-yy。
    for day_index, date in enumerate(dates):
        row = day_index + 2
        plan_sheet.cell(row, 1, date.to_pydatetime())
        plan_energy = grid_plan[day_index] * DT
        for slot, value in enumerate(plan_energy, start=2):
            plan_sheet.cell(row, slot, float(value))
        plan_sheet.cell(row, 146, float(plan_energy.sum()))
        plan_sheet.cell(row, 147, float(np.dot(price, plan_energy)))

    # 2) 充放电量：先快照模板6个数据块行的完整样式，再扩展到全部334天。
    storage_styles = [snapshot(storage_sheet, r, 6) for r in range(2, 8)]
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    four_hour_labels = [
        "0:00-4:00", "4:00-8:00", "8:00-12:00",
        "12:00-16:00", "16:00-20:00", "20:00-24:00",
    ]
    for day_index, date in enumerate(dates):
        for block, block_label in enumerate(four_hour_labels):
            row = 2 + day_index * 6 + block
            apply(storage_sheet, row, storage_styles[block])
            start = block * 24
            stop = start + 24
            storage_sheet.cell(row, 1, date.to_pydatetime() if block == 0 else None)
            storage_sheet.cell(row, 2, block_label)
            storage_sheet.cell(row, 3, float(actual_charge[day_index, start:stop].sum() * DT))
            storage_sheet.cell(row, 4, float(actual_discharge[day_index, start:stop].sum() * DT))
            if block == 0:
                storage_sheet.cell(row, 5, dt_time(0, 0))
                storage_sheet.cell(row, 6, float(soc[day_index, 0]))
            elif block == 1:
                storage_sheet.cell(row, 5, "24:00")
                storage_sheet.cell(row, 6, float(soc[day_index, -1]))

    # 3) 紧急购电量：日期只在每天第一条记录出现，其余留空，与模板一致。
    emg_first = snapshot(emergency_sheet, 2, 3)
    emg_cont = snapshot(emergency_sheet, 3, 3)
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    output_row = 2
    for day_index, date in enumerate(dates):
        nonzero_slots = np.flatnonzero(emergency[day_index] > TOL)
        for order, slot in enumerate(nonzero_slots):
            apply(emergency_sheet, output_row, emg_first if order == 0 else emg_cont)
            emergency_sheet.cell(output_row, 1, date.to_pydatetime() if order == 0 else None)
            emergency_sheet.cell(output_row, 2, slot_labels[slot])
            emergency_sheet.cell(output_row, 3, float(emergency[day_index, slot] * DT))
            output_row += 1
    if output_row == 2:
        emergency_sheet.cell(2, 1, "报告期内无紧急购电")

    workbook.save(work)
    try:
        os.replace(work, target)
        print(f"已写入 {target}")
    except PermissionError:
        fallback = target.with_name(target.stem + "_修正" + target.suffix)
        if fallback.exists():
            fallback.unlink()
        os.replace(work, fallback)
        print(f"[警告] {target.name} 被占用，已输出为 {fallback.name}；关闭占用程序后可重命名覆盖。")


def validate_results(
    price: np.ndarray,
    load_forecast: np.ndarray,
    pv_forecast: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
    grid_plan: np.ndarray,
    plan_charge: np.ndarray,
    plan_discharge: np.ndarray,
    plan_pv_used: np.ndarray,
    plan_soc: np.ndarray,
    actual: list[IntradayResult],
) -> dict[str, float | bool]:
    """独立重算关键等式和上下限，用数值证据证明闭环可行。"""

    actual_charge = np.vstack([item.charge for item in actual])
    actual_discharge = np.vstack([item.discharge for item in actual])
    emergency = np.vstack([item.emergency for item in actual])
    curtail = np.vstack([item.curtail for item in actual])
    grid_spill = np.vstack([item.grid_spill for item in actual])
    actual_soc = np.vstack([item.soc for item in actual])

    plan_balance = grid_plan + plan_discharge + plan_pv_used - load_forecast - plan_charge
    actual_balance = (
        grid_plan + pv_actual + actual_discharge + emergency
        - load_actual - actual_charge - curtail - grid_spill
    )
    actual_soc_rhs = (
        actual_soc[:, :-1]
        + ETA_C * actual_charge * DT
        - actual_discharge * DT / ETA_D
    )
    true_pv_surplus = np.maximum(pv_actual - load_actual, 0.0)

    checks: dict[str, float | bool] = {
        "SCIP可用": "SCIP" in cp.installed_solvers(),
        "日前最大功率平衡残差_kW": float(np.max(np.abs(plan_balance))),
        "日内最大功率平衡残差_kW": float(np.max(np.abs(actual_balance))),
        "实际SOC递推最大残差_kWh": float(np.max(np.abs(actual_soc[:, 1:] - actual_soc_rhs))),
        "计划SOC最小值_kWh": float(plan_soc.min()),
        "计划SOC最大值_kWh": float(plan_soc.max()),
        "实际SOC最小值_kWh": float(actual_soc.min()),
        "实际SOC最大值_kWh": float(actual_soc.max()),
        "计划终端偏差最大值_kWh": float(np.max(np.abs(plan_soc[:, -1] - PLAN_TERMINAL_SOC))),
        "跨日SOC衔接偏差最大值_kWh": float(
            np.max(np.abs(actual_soc[:-1, -1] - actual_soc[1:, 0])) if len(actual_soc) > 1 else 0.0
        ),
        "计划同时充放电最大乘积": float(np.max(plan_charge * plan_discharge)),
        "实际同时充放电最大乘积": float(np.max(actual_charge * actual_discharge)),
        "计划最大充电功率_kW": float(plan_charge.max()),
        "计划最大放电功率_kW": float(plan_discharge.max()),
        "实际最大充电功率_kW": float(actual_charge.max()),
        "实际最大放电功率_kW": float(actual_discharge.max()),
        "弃光超过真实光伏盈余最大值_kW": float(np.max(curtail - true_pv_surplus)),
        "计划购电费_元": float(np.sum(grid_plan * DT * price[None, :])),
        "紧急购电费_元": float(np.sum(emergency * DT * price[None, :] * EMERGENCY_PRICE_FACTOR)),
    }
    checks["全部硬约束通过"] = bool(
        checks["日前最大功率平衡残差_kW"] <= 1e-3
        and checks["日内最大功率平衡残差_kW"] <= 1e-3
        and checks["实际SOC递推最大残差_kWh"] <= 1e-3
        and checks["计划SOC最小值_kWh"] >= E_MIN - 1e-3
        and checks["计划SOC最大值_kWh"] <= E_MAX + 1e-3
        and checks["实际SOC最小值_kWh"] >= E_MIN - 1e-3
        and checks["实际SOC最大值_kWh"] <= E_MAX + 1e-3
        and checks["计划终端偏差最大值_kWh"] <= 1e-3
        and checks["跨日SOC衔接偏差最大值_kWh"] <= 1e-3
        and checks["计划同时充放电最大乘积"] <= 1e-3
        and checks["实际同时充放电最大乘积"] <= 1e-3
        and checks["计划最大充电功率_kW"] <= P_MAX + 1e-3
        and checks["计划最大放电功率_kW"] <= P_MAX + 1e-3
        and checks["实际最大充电功率_kW"] <= P_MAX + 1e-3
        and checks["实际最大放电功率_kW"] <= P_MAX + 1e-3
        and checks["弃光超过真实光伏盈余最大值_kW"] <= 1e-3
    )
    return checks


def make_daily_cost_plot(daily: pd.DataFrame, output_path: Path) -> None:
    """绘制计划购电费、紧急购电费和总费用的日序列。"""

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial"]
    plt.rcParams["axes.unicode_minus"] = False
    figure, axis = plt.subplots(figsize=(14, 6))
    axis.plot(daily["日期"], daily["计划购电费_元"], label="计划购电费", linewidth=1.0)
    axis.plot(daily["日期"], daily["紧急购电费_元"], label="紧急购电费", linewidth=1.0)
    axis.plot(daily["日期"], daily["实际总花费_元"], label="实际总花费", linewidth=1.4)
    axis.set_title("负荷3天/光伏5天：2025年2月1日至12月31日微网每日实际花费")
    axis.set_xlabel("日期")
    axis.set_ylabel("费用（元）")
    axis.grid(alpha=0.25)
    axis.legend()
    figure.tight_layout()
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def write_markdown_report(
    output_path: Path,
    daily: pd.DataFrame,
    checks: dict[str, float | bool],
    actual: list[IntradayResult],
) -> None:
    """输出便于论文引用的费用汇总、可行性验证和建模口径说明。"""

    actual_soc = np.vstack([item.soc for item in actual])
    lines = [
        "# 滚动预测—日前优化—日内执行运行报告（负荷3天/光伏5天）",
        "",
        "## 时间与信息口径",
        "",
        "附件时点严格按附件5顺序映射：0:10对应0:10-0:20，最后的0:00+1对应次日0:00-0:10。",
        "每天0:00只使用此前已经发生的实际数据；预测日真实负荷和真实光伏不进入日前MILP。",
        "预测窗口：负荷取同星期组最近3个历史日，光伏取最近5个自然日。",
        "",
        "## 费用结果",
        "",
        f"- 报告天数：{len(daily)}",
        f"- 计划购电费：{daily['计划购电费_元'].sum():,.2f} 元",
        f"- 紧急购电费：{daily['紧急购电费_元'].sum():,.2f} 元",
        f"- 实际总花费：{daily['实际总花费_元'].sum():,.2f} 元",
        f"- 紧急购电总量：{daily['紧急购电量_kWh'].sum():,.2f} kWh",
        f"- 弃光总量：{daily['弃光量_kWh'].sum():,.2f} kWh",
        f"- 计划购电剩余总量：{daily['计划购电剩余_kWh'].sum():,.2f} kWh",
        f"- 最终实际SOC：{actual_soc[-1, -1]:,.2f} kWh",
        "",
        "## 可行性验证",
        "",
    ]
    for name, value in checks.items():
        if isinstance(value, bool):
            shown = "通过" if value else "未通过"
        else:
            shown = f"{value:.10g}"
        lines.append(f"- {name}：{shown}")
    lines.extend([
        "",
        "## 充放电是否可以实时调整",
        "",
        "可以。题目锁定的是日前正常购电计划，不等于把储能功率也完全锁死。若储能也按预测计划刚性执行，预测误差会直接扩大紧急购电或造成无法平衡。本文把储能视为日内平衡资源：正常购电始终等于日前计划，真实数据揭示后在SOC、效率、功率和互斥约束内实时调整充放电，无法覆盖的真实缺口才使用紧急购电。该解释同时满足计划刚性和系统实时功率平衡。",
        "",
        "日前终端约束E(144)=6000约束的是预测场景下的计划轨迹。实际轨迹受预测误差影响，不强制每天回到6000，而是把当天实际日末SOC传到次日；否则跨日闭环更新会退化为每天独立运行。",
    ])
    output_path.write_text("\n".join(lines), encoding="utf-8")


def calculate_baseline_forecast_mae(
    all_dates: pd.DatetimeIndex,
    all_load: np.ndarray,
    all_pv: np.ndarray,
) -> tuple[float, float]:
    """独立重放基线最近7自然日预测，供方案D报告做同口径比较。"""

    load_errors, pv_errors = [], []
    date_to_index = {date: index for index, date in enumerate(all_dates)}
    for date in pd.date_range(START_DATE, END_DATE, freq="D"):
        index = date_to_index[date]
        load_forecast = moving_average_3(np.median(all_load[index - 7:index], axis=0))
        pv_forecast = moving_average_3(np.median(all_pv[index - 7:index], axis=0))
        load_errors.append(np.mean(np.abs(load_forecast - all_load[index])))
        pv_errors.append(np.mean(np.abs(pv_forecast - all_pv[index])))
    return float(np.mean(load_errors)), float(np.mean(pv_errors))


def write_scheme_d_report(
    output_path: Path,
    daily: pd.DataFrame,
    checks: dict[str, float | bool],
    baseline_load_mae: float,
    baseline_pv_mae: float,
) -> None:
    """生成方案D与最开始基线模型的定量对比报告。"""

    baseline_path = SCRIPT_DIR.parent / "每日费用.csv"
    baseline = pd.read_csv(baseline_path) if baseline_path.exists() else None
    plan_cost = float(daily["计划购电费_元"].sum())
    emergency_cost = float(daily["紧急购电费_元"].sum())
    total_cost = float(daily["实际总花费_元"].sum())
    lines = [
        "# 方案D运行与对比报告",
        "",
        "## 方法",
        "",
        "负荷预测按星期分为两组：周五、周六为一组；周一、周二、周三、周四、周日为另一组。预测目标日时，只使用同组最近7个历史日的同一十分钟时点中位数，再做三点移动平滑。光伏仍使用最近7个自然日同一时点中位数和三点平滑。",
        "",
        "规划模型、日内执行和基线保持一致，以便只检验星期分组预测的贡献。充电效率与放电效率均为0.9，SOC递推为 `E(t+1)=E(t)+0.9*P_charge*Δt-P_discharge*Δt/0.9`。",
        "",
        "## 方案D结果",
        "",
        f"- 计划购电费：{plan_cost:,.2f} 元",
        f"- 紧急购电费：{emergency_cost:,.2f} 元",
        f"- 实际总花费：{total_cost:,.2f} 元",
        f"- 紧急购电量：{daily['紧急购电量_kWh'].sum():,.2f} kWh",
        f"- 负荷预测MAE：{daily['负荷预测MAE_kW'].mean():,.2f} kW",
        f"- 光伏预测MAE：{daily['光伏预测MAE_kW'].mean():,.2f} kW",
        "",
        "## 预测精度对比",
        "",
        "| 指标 | 基线最近7自然日 | 方案D | 变化率 |",
        "|---|---:|---:|---:|",
        f"| 负荷MAE（kW） | {baseline_load_mae:,.2f} | {daily['负荷预测MAE_kW'].mean():,.2f} | {(daily['负荷预测MAE_kW'].mean()/baseline_load_mae-1)*100:+.2f}% |",
        f"| 光伏MAE（kW） | {baseline_pv_mae:,.2f} | {daily['光伏预测MAE_kW'].mean():,.2f} | {(daily['光伏预测MAE_kW'].mean()/baseline_pv_mae-1)*100:+.2f}% |",
        "",
    ]
    if baseline is not None and len(baseline) == len(daily):
        base_plan = float(baseline["计划购电费_元"].sum())
        base_emergency = float(baseline["紧急购电费_元"].sum())
        base_total = float(baseline["实际总花费_元"].sum())
        lines.extend([
            "## 与基线对比",
            "",
            "| 指标 | 基线 | 方案D | 变化 |",
            "|---|---:|---:|---:|",
            f"| 计划购电费（元） | {base_plan:,.2f} | {plan_cost:,.2f} | {plan_cost-base_plan:+,.2f} |",
            f"| 紧急购电费（元） | {base_emergency:,.2f} | {emergency_cost:,.2f} | {emergency_cost-base_emergency:+,.2f} |",
            f"| 实际总花费（元） | {base_total:,.2f} | {total_cost:,.2f} | {total_cost-base_total:+,.2f} |",
            "",
        ])
    lines.extend([
        "## 可行性验证",
        "",
        f"- 全部硬约束：{'通过' if checks['全部硬约束通过'] else '未通过'}",
        f"- 日前最大功率平衡残差：{checks['日前最大功率平衡残差_kW']:.3e} kW",
        f"- 日内最大功率平衡残差：{checks['日内最大功率平衡残差_kW']:.3e} kW",
        f"- 实际SOC范围：{checks['实际SOC最小值_kWh']:.4f}—{checks['实际SOC最大值_kWh']:.4f} kWh",
        f"- 计划同时充放电最大乘积：{checks['计划同时充放电最大乘积']:.3e}",
        f"- 实际同时充放电最大乘积：{checks['实际同时充放电最大乘积']:.3e}",
        f"- 充电效率：{ETA_C}",
        f"- 放电效率：{ETA_D}",
    ])
    output_path.write_text("\n".join(lines), encoding="utf-8")


def write_comparison_report(
    output_path: Path,
    daily: pd.DataFrame,
    checks: dict[str, float | bool],
) -> None:
    """生成新方案(负荷3天/光伏5天)与旧方案(方案D 7/7)及最初基线的对比报告。"""

    old = pd.read_csv(OLD_SCHEME_DAILY) if OLD_SCHEME_DAILY.exists() else None
    base = pd.read_csv(BASE_SCHEME_DAILY) if BASE_SCHEME_DAILY.exists() else None

    new_plan = float(daily["计划购电费_元"].sum())
    new_emg = float(daily["紧急购电费_元"].sum())
    new_total = float(daily["实际总花费_元"].sum())
    new_emg_q = float(daily["紧急购电量_kWh"].sum())
    new_curtail = float(daily["弃光量_kWh"].sum())
    new_load_mae = float(daily["负荷预测MAE_kW"].mean())
    new_pv_mae = float(daily["光伏预测MAE_kW"].mean())

    lines = [
        "# 问题二 优化窗口方案（负荷3天 / 光伏5天）运行与对比报告",
        "",
        "## 1. 方案说明",
        "",
        "本方案相对旧方案（方案D）只改变预测窗口长度，日前优化、日内执行、储能参数完全不变：",
        "",
        "- 负荷预测：星期分组（周五/周六一组，周一~周四+周日一组），取同组最近 **3** 个历史日，逐十分钟时点取中位数，再做三点移动平滑。",
        "- 光伏预测：取最近 **5** 个自然日，逐十分钟时点取中位数，再做三点移动平滑。",
        "- 窗口依据：阶梯测试在2025-01-25~31上给出的最优组合（负荷3天、光伏5天）。",
        "- 变量单位kW；Excel中购电/充放电为十分钟电量kWh（乘 Δt=1/6 h）。",
        "",
        "## 2. 本方案结果",
        "",
        f"- 计划购电费：{new_plan:,.2f} 元",
        f"- 紧急购电费：{new_emg:,.2f} 元",
        f"- 实际总花费：{new_total:,.2f} 元",
        f"- 紧急购电量：{new_emg_q:,.2f} kWh",
        f"- 弃光量：{new_curtail:,.2f} kWh",
        f"- 负荷预测MAE：{new_load_mae:,.2f} kW",
        f"- 光伏预测MAE：{new_pv_mae:,.2f} kW",
        "",
    ]

    if old is not None:
        o_plan = float(old["计划购电费_元"].sum())
        o_emg = float(old["紧急购电费_元"].sum())
        o_total = float(old["实际总花费_元"].sum())
        o_emg_q = float(old["紧急购电量_kWh"].sum())
        o_curtail = float(old["弃光量_kWh"].sum())
        o_load_mae = float(old["负荷预测MAE_kW"].mean())
        o_pv_mae = float(old["光伏预测MAE_kW"].mean())
        lines += [
            "## 3. 与旧方案（方案D：负荷7天/光伏7天）对比",
            "",
            "| 指标 | 旧方案(7/7) | 新方案(3/5) | 变化 | 变化率 |",
            "|---|---:|---:|---:|---:|",
            f"| 计划购电费（元） | {o_plan:,.2f} | {new_plan:,.2f} | {new_plan-o_plan:+,.2f} | {(new_plan/o_plan-1)*100:+.2f}% |",
            f"| 紧急购电费（元） | {o_emg:,.2f} | {new_emg:,.2f} | {new_emg-o_emg:+,.2f} | {(new_emg/o_emg-1)*100:+.2f}% |",
            f"| 实际总花费（元） | {o_total:,.2f} | {new_total:,.2f} | {new_total-o_total:+,.2f} | {(new_total/o_total-1)*100:+.2f}% |",
            f"| 紧急购电量（kWh） | {o_emg_q:,.2f} | {new_emg_q:,.2f} | {new_emg_q-o_emg_q:+,.2f} | {(new_emg_q/o_emg_q-1)*100:+.2f}% |",
            f"| 弃光量（kWh） | {o_curtail:,.2f} | {new_curtail:,.2f} | {new_curtail-o_curtail:+,.2f} | {(new_curtail/o_curtail-1)*100:+.2f}% |",
            f"| 负荷预测MAE（kW） | {o_load_mae:,.2f} | {new_load_mae:,.2f} | {new_load_mae-o_load_mae:+,.2f} | {(new_load_mae/o_load_mae-1)*100:+.2f}% |",
            f"| 光伏预测MAE（kW） | {o_pv_mae:,.2f} | {new_pv_mae:,.2f} | {new_pv_mae-o_pv_mae:+,.2f} | {(new_pv_mae/o_pv_mae-1)*100:+.2f}% |",
            "",
            f"**相对旧方案总费用变化：{new_total-o_total:+,.2f} 元（{(new_total/o_total-1)*100:+.2f}%）。**",
            "",
        ]

    if base is not None:
        b_plan = float(base["计划购电费_元"].sum())
        b_emg = float(base["紧急购电费_元"].sum())
        b_total = float(base["实际总花费_元"].sum())
        lines += [
            "## 4. 与最初基线（负荷7自然日/光伏7自然日、负荷不分组）对比",
            "",
            "| 指标 | 最初基线 | 新方案(3/5) | 变化 | 变化率 |",
            "|---|---:|---:|---:|---:|",
            f"| 计划购电费（元） | {b_plan:,.2f} | {new_plan:,.2f} | {new_plan-b_plan:+,.2f} | {(new_plan/b_plan-1)*100:+.2f}% |",
            f"| 紧急购电费（元） | {b_emg:,.2f} | {new_emg:,.2f} | {new_emg-b_emg:+,.2f} | {(new_emg/b_emg-1)*100:+.2f}% |",
            f"| 实际总花费（元） | {b_total:,.2f} | {new_total:,.2f} | {new_total-b_total:+,.2f} | {(new_total/b_total-1)*100:+.2f}% |",
            "",
        ]

    lines += [
        "## 5. 可行性验证",
        "",
        f"- 全部硬约束：{'通过' if checks['全部硬约束通过'] else '未通过'}",
        f"- 日前最大功率平衡残差：{checks['日前最大功率平衡残差_kW']:.3e} kW",
        f"- 日内最大功率平衡残差：{checks['日内最大功率平衡残差_kW']:.3e} kW",
        f"- 实际SOC最小值/最大值：{checks['实际SOC最小值_kWh']:.4f} / {checks['实际SOC最大值_kWh']:.4f} kWh",
        f"- 计划同时充放电最大乘积：{checks['计划同时充放电最大乘积']:.3e}",
        f"- 实际同时充放电最大乘积：{checks['实际同时充放电最大乘积']:.3e}",
        f"- 充电效率：{ETA_C}；放电效率：{ETA_D}",
        "",
        "## 6. 说明",
        "",
        "本方案与旧方案仅预测窗口不同，费用差异可归因于预测精度变化。",
    ]
    output_path.write_text("\n".join(lines), encoding="utf-8")


def run_pseudo_test() -> None:
    """用合成数据快速检查预测、MILP、日内平衡和边界约束。"""

    slots = np.arange(N_SLOTS)
    price = 0.45 + 0.25 * ((slots >= 102) & (slots < 126))
    history_load = np.vstack([
        4000 + 600 * np.sin(2 * np.pi * (slots - 30) / N_SLOTS) + day * 4
        for day in range(31)
    ])
    daylight = np.maximum(0.0, np.sin(np.pi * (slots - 36) / 72))
    history_pv = np.vstack([5000 * daylight * (0.9 + day / 1000) for day in range(31)])
    load_forecast, pv_forecast = forecast_from_history(history_load, history_pv)
    problem, objects = build_day_ahead_problem(price)
    plan = solve_day_ahead(problem, objects, load_forecast, pv_forecast, E_INIT)
    actual_load = load_forecast * (1.0 + 0.03 * np.sin(slots / 5))
    actual_pv = pv_forecast * (1.0 - 0.08 * np.cos(slots / 7))
    actual = execute_intraday(plan.grid, actual_load, actual_pv, E_INIT)
    balance = (
        plan.grid + actual_pv + actual.discharge + actual.emergency
        - actual_load - actual.charge - actual.curtail - actual.grid_spill
    )
    assert np.max(np.abs(balance)) < 1e-6
    assert actual.soc.min() >= E_MIN - TOL and actual.soc.max() <= E_MAX + TOL
    assert np.max(actual.charge * actual.discharge) <= TOL
    print("伪数据测试通过：预测、SCIP日前MILP、固定购电日内执行和SOC递推均正常。")


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="运行问题二微网调度闭环")
    parser.add_argument("--pseudo-test", action="store_true", help="仅运行合成数据逻辑测试")
    parser.add_argument("--end-date", default="2025-12-31", help="调试时可缩短运行截止日")
    return parser.parse_args()


def main() -> None:
    arguments = parse_arguments()
    if arguments.pseudo_test:
        run_pseudo_test()
        return

    price, all_dates, all_load, all_pv, slot_labels = read_input_data()
    if ETA_C != 0.9 or ETA_D != 0.9:
        raise AssertionError("题目要求 ETA_C = ETA_D = 0.9。")
    requested_end = pd.Timestamp(arguments.end_date)
    if requested_end < START_DATE or requested_end > END_DATE:
        raise ValueError("--end-date 必须位于 2025-02-01 至 2025-12-31。")
    report_dates = pd.date_range(START_DATE, requested_end, freq="D")
    date_to_index = {date: index for index, date in enumerate(all_dates)}

    # 1月全部实际数据作为预热训练集；2月起每天闭环追加实际值。
    january_mask = all_dates < START_DATE
    load_history = [row.copy() for row in all_load[january_mask]]
    pv_history = [row.copy() for row in all_pv[january_mask]]
    history_dates = list(all_dates[january_mask])

    problem, objects = build_day_ahead_problem(price)
    n_days = len(report_dates)
    load_forecasts = np.zeros((n_days, N_SLOTS))
    pv_forecasts = np.zeros((n_days, N_SLOTS))
    grid_plans = np.zeros((n_days, N_SLOTS))
    plan_charges = np.zeros((n_days, N_SLOTS))
    plan_discharges = np.zeros((n_days, N_SLOTS))
    plan_pv_used = np.zeros((n_days, N_SLOTS))
    plan_socs = np.zeros((n_days, N_SLOTS + 1))
    actual_results: list[IntradayResult] = []
    actual_soc_start = E_INIT
    daily_rows: list[dict[str, float | pd.Timestamp]] = []
    detail_frames: list[pd.DataFrame] = []

    for day_number, date in enumerate(report_dates):
        source_index = date_to_index[date]
        load_forecast, pv_forecast = forecast_from_history(
            np.asarray(load_history),
            np.asarray(pv_history),
            target_date=date,
            history_dates=pd.DatetimeIndex(history_dates),
        )
        plan = solve_day_ahead(
            problem, objects, load_forecast, pv_forecast, actual_soc_start
        )
        actual = execute_intraday(
            plan.grid, all_load[source_index], all_pv[source_index], actual_soc_start
        )

        load_forecasts[day_number] = load_forecast
        pv_forecasts[day_number] = pv_forecast
        grid_plans[day_number] = plan.grid
        plan_charges[day_number] = plan.charge
        plan_discharges[day_number] = plan.discharge
        plan_pv_used[day_number] = plan.pv_used
        plan_socs[day_number] = plan.soc
        actual_results.append(actual)

        plan_cost = float(np.dot(price, plan.grid * DT))
        emergency_cost = float(np.dot(price * EMERGENCY_PRICE_FACTOR, actual.emergency * DT))
        daily_rows.append({
            "日期": date,
            "计划购电量_kWh": float(plan.grid.sum() * DT),
            "计划购电费_元": plan_cost,
            "紧急购电量_kWh": float(actual.emergency.sum() * DT),
            "紧急购电费_元": emergency_cost,
            "实际总花费_元": plan_cost + emergency_cost,
            "实际充电量_kWh": float(actual.charge.sum() * DT),
            "实际放电量_kWh": float(actual.discharge.sum() * DT),
            "弃光量_kWh": float(actual.curtail.sum() * DT),
            "计划购电剩余_kWh": float(actual.grid_spill.sum() * DT),
            "日初SOC_kWh": float(actual.soc[0]),
            "日末SOC_kWh": float(actual.soc[-1]),
            "负荷预测MAE_kW": float(np.mean(np.abs(load_forecast - all_load[source_index]))),
            "光伏预测MAE_kW": float(np.mean(np.abs(pv_forecast - all_pv[source_index]))),
        })

        if date in SPECIAL_DATES:
            detail_frames.append(pd.DataFrame({
                "日期": date,
                "时段": slot_labels,
                "电价_元每kWh": price,
                "预测负荷_kW": load_forecast,
                "实际负荷_kW": all_load[source_index],
                "预测光伏_kW": pv_forecast,
                "实际光伏_kW": all_pv[source_index],
                "计划购电_kW": plan.grid,
                "计划充电_kW": plan.charge,
                "计划放电_kW": plan.discharge,
                "实际充电_kW": actual.charge,
                "实际放电_kW": actual.discharge,
                "紧急购电_kW": actual.emergency,
                "弃光_kW": actual.curtail,
                "计划购电剩余_kW": actual.grid_spill,
                "时段末SOC_kWh": actual.soc[1:],
            }))

        # 闭环更新：当天结束后才把当天真实数据放入历史库。
        load_history.append(all_load[source_index].copy())
        pv_history.append(all_pv[source_index].copy())
        history_dates.append(date)
        actual_soc_start = float(actual.soc[-1])

        if day_number == 0 or (day_number + 1) % 25 == 0 or day_number + 1 == n_days:
            print(
                f"进度 {day_number + 1:3d}/{n_days}: {date.date()}，"
                f"累计实际费用 {sum(row['实际总花费_元'] for row in daily_rows):,.2f} 元"
            )

    actual_charge = np.vstack([item.charge for item in actual_results])
    actual_discharge = np.vstack([item.discharge for item in actual_results])
    emergency = np.vstack([item.emergency for item in actual_results])
    actual_soc = np.vstack([item.soc for item in actual_results])
    actual_load_report = np.vstack([all_load[date_to_index[date]] for date in report_dates])
    actual_pv_report = np.vstack([all_pv[date_to_index[date]] for date in report_dates])

    checks = validate_results(
        price,
        load_forecasts,
        pv_forecasts,
        actual_load_report,
        actual_pv_report,
        grid_plans,
        plan_charges,
        plan_discharges,
        plan_pv_used,
        plan_socs,
        actual_results,
    )
    if not checks["全部硬约束通过"]:
        raise AssertionError(f"可行性验证未通过：{checks}")

    daily = pd.DataFrame(daily_rows)
    result_path = SCRIPT_DIR / "result2.xlsx"
    write_result_workbook(
        result_path,
        report_dates,
        slot_labels,
        price,
        grid_plans,
        actual_charge,
        actual_discharge,
        emergency,
        actual_soc,
    )
    daily.to_csv(SCRIPT_DIR / "每日费用.csv", index=False, encoding="utf-8-sig")
    make_daily_cost_plot(daily, SCRIPT_DIR / "费用对比.png")
    write_markdown_report(SCRIPT_DIR / "运行报告.md", daily, checks, actual_results)
    write_comparison_report(SCRIPT_DIR / "对比报告.md", daily, checks)

    print("\n=== 全年费用汇总 ===")
    print(daily.to_string(index=False, max_rows=20))
    print(f"\n计划购电费：{daily['计划购电费_元'].sum():,.2f} 元")
    print(f"紧急购电费：{daily['紧急购电费_元'].sum():,.2f} 元")
    print(f"实际总花费：{daily['实际总花费_元'].sum():,.2f} 元")
    print("\n=== 可行性验证 ===")
    for name, value in checks.items():
        print(f"{name}: {value}")

    print("\n=== 四个指定日期的144时段详细表 ===")
    if detail_frames:
        detailed = pd.concat(detail_frames, ignore_index=True)
        detailed.to_csv(SCRIPT_DIR / "四日详细表.csv", index=False, encoding="utf-8-sig")
        with pd.option_context("display.max_rows", None, "display.max_columns", None, "display.width", 260):
            print(detailed.to_string(index=False))
    else:
        print("当前 --end-date 尚未覆盖指定日期。")

    print(f"\n结果文件：{result_path}")


if __name__ == "__main__":
    main()
